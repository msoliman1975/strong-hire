"""Test helpers for the input jobs, shared by the worker and API test suites.

Nothing here runs in production. It lives in the package so both test folders can import it.
"""

from __future__ import annotations

import io
import json
from collections.abc import Callable, Sequence
from pathlib import Path
from typing import Any

import docx
import httpx
from pydantic import BaseModel
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles

from strong_core.gateway import FakeBackend, Message, Role, TokenUsage
from strong_core.gateway.fake import DEFAULT_FIXTURES_DIR
from strong_worker.inputs.fetch import PostingFetcher


def enable_sqlite_jsonb() -> None:
    """Let the Postgres models create tables on SQLite (JSONB becomes JSON) for unit tests."""

    @compiles(JSONB, "sqlite")
    def _jsonb_on_sqlite(type_: Any, compiler: Any, **kw: Any) -> str:
        return "JSON"


class RecordingBackend(FakeBackend):
    """Fake model backend that keeps every message list it gets, and can fail on purpose."""

    def __init__(self, fixtures_dir: Path, fail_times: int = 0) -> None:
        super().__init__(fixtures_dir)
        self.calls: list[list[Message]] = []
        self.fail_times = fail_times

    def complete(
        self, role: Role, messages: Sequence[Message], output_type: type[BaseModel] | None
    ) -> tuple[object, TokenUsage]:
        self.calls.append(list(messages))
        if self.fail_times > 0:
            self.fail_times -= 1
            raise RuntimeError("simulated model failure")
        return super().complete(role, messages, output_type)


def copy_fake_fixtures(target: Path) -> Path:
    for src in DEFAULT_FIXTURES_DIR.rglob("*"):
        if src.is_file():
            dst = target / src.relative_to(DEFAULT_FIXTURES_DIR)
            dst.parent.mkdir(parents=True, exist_ok=True)
            dst.write_bytes(src.read_bytes())
    return target


def set_extractor_output(fixtures: Path, output_type: str, data: dict[str, Any]) -> None:
    """Make the fake extractor return `data` for every `output_type` request."""
    (fixtures / "extractor" / f"{output_type}.json").write_text(json.dumps(data), encoding="utf-8")


async def public_resolver(host: str) -> list[str]:
    return ["93.184.216.34"]


Route = httpx.Response | Callable[[httpx.Request], httpx.Response]


def make_fetcher(
    routes: dict[str, Route],
    *,
    rendered: str | None = None,
    resolver: Callable[[str], Any] = public_resolver,
) -> PostingFetcher:
    """A fetcher over a fake network. Routes map 'host/path' to a response or a handler."""

    def handler(request: httpx.Request) -> httpx.Response:
        route = routes.get(f"{request.url.host}{request.url.path}")
        if route is None:
            return httpx.Response(404, text="not found")
        return route(request) if callable(route) else route

    async def renderer(url: str) -> str | None:
        return rendered

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    return PostingFetcher(client, resolver=resolver, renderer=renderer)


def make_pdf(lines: list[str]) -> bytes:
    """A one-page PDF with a real text layer, built by hand."""

    def esc(s: str) -> str:
        s = s.replace("\u2022", "-").encode("latin-1", "replace").decode("latin-1")
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    content = "BT /F1 11 Tf 50 760 Td 14 TL " + " ".join(f"({esc(x)}) Tj T*" for x in lines)
    content += " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_docx(paragraphs: list[str], hidden: list[str] | None = None) -> bytes:
    """A DOCX file. `hidden` paragraphs use hidden font runs, which a reader cannot see."""
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    for text in hidden or []:
        document.add_paragraph().add_run(text).font.hidden = True
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
