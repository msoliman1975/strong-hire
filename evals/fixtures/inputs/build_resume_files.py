"""Build the synthetic PDF and DOCX resume fixtures in resumes/ and their expected JSON.

    uv run python evals/fixtures/inputs/build_resume_files.py

Every person, employer and school here is made up. Each case is a layout that a text reader can
get wrong: a date column on the left, a sidebar next to the main column, bold sub-sections,
a table, a role split across two pages, two roles at one company, a DOCX table and a DOCX text
box. The expected JSON is written from the same data, so the file and the answer always agree.

Some PDFs draw their text in an order that is not the reading order, as design tools and
word processors with floating text boxes do. `date-column` copies a real failure: the next
role's heading is drawn before the sub-sections of the role above it.

The output is stable: reportlab runs in invariant mode and DOCX zip entries get a fixed date.
Needs the dev dependencies (reportlab, python-docx).
"""

from __future__ import annotations

import io
import json
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import docx
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen.canvas import Canvas

OUT = Path(__file__).parent / "resumes"
FIXED_ZIP_DATE = (2026, 1, 1, 0, 0, 0)
REGULAR, BOLD = "Helvetica", "Helvetica-Bold"


@dataclass
class Sub:
    """A bold lead-in ("Search Platform:") and its text, on one line or wrapped."""

    head: str
    text: str

    @property
    def achievement(self) -> str:
        return f"{self.head} {self.text}"


@dataclass
class Role:
    company: str
    place: str
    title: str
    dates: str  # as printed
    start: str | None
    end: str | None
    summary: str | None = None  # a short paragraph under the title; expected as an achievement
    subs: list[Sub] = field(default_factory=list)
    bullets: list[str] = field(default_factory=list)
    groups: list[tuple[str, list[str]]] = field(default_factory=list)  # bold heading + bullets
    skills: list[str] = field(default_factory=list)

    def achievements(self) -> list[str]:
        items = [self.summary] if self.summary else []
        items += [s.achievement for s in self.subs]
        items += self.bullets
        for _, bullets in self.groups:
            items += bullets
        return items


@dataclass
class Case:
    name: str
    person: str
    contact: str
    summary: str | None
    roles: list[Role]
    skills: list[str]
    education: list[dict[str, Any]]
    certifications: list[str] = field(default_factory=list)

    def expected(self) -> dict[str, Any]:
        return {
            "summary": self.summary,
            "roles": [
                {
                    "title": r.title,
                    "company": r.company,
                    "start": r.start,
                    "end": r.end,
                    "achievements": r.achievements(),
                    "skills": r.skills,
                }
                for r in self.roles
            ],
            "skills": self.skills,
            "education": self.education,
            "certifications": self.certifications,
        }

    def education_lines(self) -> list[str]:
        out = []
        for e in self.education:
            degree = " ".join(x for x in (e.get("degree"), e.get("field_of_study")) if x)
            out.append(f"{degree}, {e['institution']}, {e['end']}" if degree else e["institution"])
        return out


def edu(institution: str, degree: str, field_of_study: str | None, end: str) -> dict[str, Any]:
    return {
        "institution": institution,
        "degree": degree,
        "field_of_study": field_of_study,
        "end": end,
    }


# --- PDF drawing --------------------------------------------------------------------------------

Op = tuple[int, float, float, str, str, float]  # page, x, y, font, text, size


class Layout:
    """Lines placed top-down in reading order. Each op gets a draw key; lower keys draw first."""

    def __init__(self, top: float = 740, bottom: float = 60) -> None:
        self.ops: list[tuple[tuple[int, ...], Op]] = []
        self.page = 0
        self.top = top
        self.bottom = bottom
        self.y = top
        self.seq = 0

    def put(
        self,
        key: tuple[int, ...],
        x: float,
        text: str,
        font: str = REGULAR,
        size: float = 10,
        y: float | None = None,
    ) -> None:
        self.seq += 1
        self.ops.append(
            ((*key, self.seq), (self.page, x, self.y if y is None else y, font, text, size))
        )

    def down(self, by: float = 13) -> None:
        self.y -= by
        if self.y < self.bottom:
            self.new_page()

    def new_page(self) -> None:
        self.page += 1
        self.y = self.top

    def wrap(
        self,
        key: tuple[int, ...],
        x: float,
        width: float,
        text: str,
        size: float = 10,
        lead: str | None = None,
        indent: float = 0,
    ) -> None:
        """Wrapped text; `lead` is a bold lead-in on the first line."""
        first = True
        prefix_w = 0.0
        if lead:
            prefix_w = _width(lead + " ", BOLD, size)
        for line in _split(text, width - prefix_w if lead else width - indent, size):
            if first and lead:
                self.put(key, x, lead, BOLD, size)
                self.put(key, x + prefix_w, line, REGULAR, size)
            else:
                self.put(key, x + (0 if first else indent), line, REGULAR, size)
            first = False
            self.down()


def _width(text: str, font: str, size: float) -> float:
    from reportlab.pdfbase.pdfmetrics import stringWidth

    return float(stringWidth(text, font, size))


def _split(text: str, width: float, size: float) -> list[str]:
    return [str(x) for x in simpleSplit(text, REGULAR, size, width)]


def render_pdf(
    layout: Layout, footer: Callable[[int, int], str] | None = None, header: str | None = None
) -> bytes:
    buf = io.BytesIO()
    canvas = Canvas(buf, pagesize=LETTER, invariant=1)
    canvas.setTitle("Resume")
    canvas.setAuthor("Strong Hire fixtures")
    pages = layout.page + 1
    for page in range(pages):
        if header and page > 0:
            canvas.setFont(REGULAR, 8)
            canvas.drawString(50, 770, header)
        for _, (p, x, y, font, text, size) in sorted(layout.ops, key=lambda o: o[0]):
            if p == page:
                canvas.setFont(font, size)
                canvas.drawString(x, y, text)
        if footer:
            canvas.setFont(REGULAR, 8)
            canvas.drawString(270, 30, footer(page + 1, pages))
        canvas.showPage()
    canvas.save()
    return buf.getvalue()


def head_block(layout: Layout, key: tuple[int, ...], case: Case, x: float = 50) -> None:
    layout.put(key, x, case.person, BOLD, 16)
    layout.down(18)
    layout.put(key, x, case.contact, REGULAR, 9)
    layout.down(20)
    if case.summary:
        layout.put(key, x, "SUMMARY", BOLD, 11)
        layout.down(14)
        layout.wrap(key, x, 510, case.summary)
        layout.down(6)


def tail_block(
    layout: Layout, key: tuple[int, ...], case: Case, x: float = 50, width: float = 510
) -> None:
    layout.down(6)
    layout.put(key, x, "SKILLS", BOLD, 11)
    layout.down(14)
    layout.wrap(key, x, width, ", ".join(case.skills))
    layout.down(6)
    layout.put(key, x, "EDUCATION", BOLD, 11)
    layout.down(14)
    for line in case.education_lines():
        layout.wrap(key, x, width, line)
    if case.certifications:
        layout.down(6)
        layout.put(key, x, "CERTIFICATIONS", BOLD, 11)
        layout.down(14)
        for cert in case.certifications:
            layout.wrap(key, x, width, cert)


def role_body(layout: Layout, key: tuple[int, ...], role: Role, x: float, width: float) -> None:
    for sub in role.subs:
        layout.wrap(key, x, width, sub.text, lead=sub.head)
    for bullet in role.bullets:
        layout.wrap(key, x, width, "- " + bullet, indent=8)
    for heading, bullets in role.groups:
        layout.put(key, x, heading, BOLD, 10)
        layout.down()
        for bullet in bullets:
            layout.wrap(key, x + 8, width - 8, "- " + bullet, indent=8)
    if role.skills:
        layout.wrap(key, x, width, "Tools: " + ", ".join(role.skills))
    layout.down(8)


def pdf_date_column(case: Case) -> bytes:
    """Dates in a left column; heading, title and summary of the next role draw before the
    sub-sections of the role above it (the order of the failure seen on the test server)."""
    layout = Layout()
    head_block(layout, (0,), case)
    layout.put((0,), 50, "EXPERIENCE", BOLD, 11)
    layout.down(16)
    for i, role in enumerate(case.roles):
        head_key = (1, 0) if i == 0 else (1, 2 * i - 1)
        layout.put(head_key, 50, role.dates, REGULAR, 9)
        layout.put(head_key, 150, f"{role.company}  {role.place}", BOLD, 10)
        layout.down()
        layout.put(head_key, 150, role.title, REGULAR, 10)
        layout.down()
        if role.summary:
            layout.wrap(head_key, 150, 410, role.summary)
        role_body(layout, (1, 2 * i + 2), role, 150, 410)
    tail_block(layout, (9,), case)
    return render_pdf(layout)


def pdf_sidebar(case: Case) -> bytes:
    """A 150 pt sidebar (contact, skills, education) next to the main column."""
    layout = Layout(top=750)
    layout.put((0,), 40, case.person, BOLD, 16)
    layout.down(26)
    columns_top = layout.y
    for part in case.contact.split(" | "):
        layout.put((0,), 40, part, REGULAR, 8)
        layout.down(11)
    layout.down(8)
    layout.put((0,), 40, "SKILLS", BOLD, 10)
    layout.down(13)
    for skill in case.skills:
        layout.put((0,), 40, skill, REGULAR, 9)
        layout.down(12)
    layout.down(8)
    layout.put((0,), 40, "EDUCATION", BOLD, 10)
    layout.down(13)
    for e in case.education:
        layout.wrap((0,), 40, 140, f"{e['degree']} {e['field_of_study'] or ''}".strip(), size=9)
        layout.wrap((0,), 40, 140, f"{e['institution']}, {e['end']}", size=9)
    if case.certifications:
        layout.down(8)
        layout.put((0,), 40, "CERTIFICATIONS", BOLD, 10)
        layout.down(13)
        for cert in case.certifications:
            layout.wrap((0,), 40, 140, cert, size=9)

    layout.y = columns_top
    if case.summary:
        layout.put((1,), 210, "PROFILE", BOLD, 11)
        layout.down(14)
        layout.wrap((1,), 210, 360, case.summary)
        layout.down(6)
    layout.put((1,), 210, "EXPERIENCE", BOLD, 11)
    layout.down(16)
    for role in case.roles:
        layout.put((1,), 210, role.title, BOLD, 10)
        layout.down()
        layout.put((1,), 210, f"{role.company}, {role.place}  |  {role.dates}", REGULAR, 9)
        layout.down()
        role_body(layout, (1,), role, 210, 360)
    return render_pdf(layout)


def pdf_single_column(
    case: Case, *, footer: bool = False, header: bool = False, top: float = 740
) -> bytes:
    """One column, title and right-aligned dates on one line, company on the next."""
    layout = Layout(top=top)
    head_block(layout, (0,), case)
    layout.put((0,), 50, "EXPERIENCE", BOLD, 11)
    layout.down(16)
    for role in case.roles:
        layout.put((0,), 50, role.title, BOLD, 10.5)
        layout.put((0,), 560 - _width(role.dates, REGULAR, 9), role.dates, REGULAR, 9)
        layout.down()
        layout.put((0,), 50, f"{role.company}, {role.place}", REGULAR, 10)
        layout.down()
        if role.summary:
            layout.wrap((0,), 50, 510, role.summary)
        role_body(layout, (0,), role, 50, 510)
    tail_block(layout, (0,), case)
    return render_pdf(
        layout,
        footer=(lambda p, n: f"Page {p} of {n}") if footer else None,
        header=f"{case.person} - Resume" if header else None,
    )


def pdf_same_company(case: Case) -> bytes:
    """The company heading once, then each title with its dates under it."""
    layout = Layout()
    head_block(layout, (0,), case)
    layout.put((0,), 50, "EXPERIENCE", BOLD, 11)
    layout.down(16)
    last = None
    for role in case.roles:
        if role.company != last:
            layout.down(4)
            layout.put((0,), 50, f"{role.company.upper()}  -  {role.place}", BOLD, 11)
            layout.down(15)
            last = role.company
        layout.put((0,), 60, role.title, BOLD, 10)
        layout.put((0,), 560 - _width(role.dates, REGULAR, 9), role.dates, REGULAR, 9)
        layout.down()
        role_body(layout, (0,), role, 60, 500)
    tail_block(layout, (0,), case)
    return render_pdf(layout)


def pdf_table(case: Case) -> bytes:
    """Experience as a ruled table: a dates cell and a details cell per row."""
    buf_layout = Layout()
    head_block(buf_layout, (0,), case)
    buf_layout.put((0,), 50, "EXPERIENCE", BOLD, 11)
    buf_layout.down(14)
    rules: list[tuple[int, float]] = []
    for role in case.roles:
        rules.append((buf_layout.page, buf_layout.y + 11))
        buf_layout.put((0,), 55, role.dates, REGULAR, 9)
        buf_layout.put((0,), 175, f"{role.title}, {role.company} ({role.place})", BOLD, 10)
        buf_layout.down()
        role_body(buf_layout, (0,), role, 175, 385)
    rules.append((buf_layout.page, buf_layout.y + 11))
    tail_block(buf_layout, (0,), case)

    buf = io.BytesIO()
    canvas = Canvas(buf, pagesize=LETTER, invariant=1)
    canvas.setTitle("Resume")
    for page in range(buf_layout.page + 1):
        ys = [y for p, y in rules if p == page]
        if ys:
            canvas.setLineWidth(0.5)
            for y in ys:
                canvas.line(50, y, 562, y)
            canvas.line(50, max(ys), 50, min(ys))
            canvas.line(168, max(ys), 168, min(ys))
            canvas.line(562, max(ys), 562, min(ys))
        for _, (p, x, y, font, text, size) in sorted(buf_layout.ops, key=lambda o: o[0]):
            if p == page:
                canvas.setFont(font, size)
                canvas.drawString(x, y, text)
        canvas.showPage()
    canvas.save()
    return buf.getvalue()


# --- DOCX -----------------------------------------------------------------------------------


def _docx_bytes(document: Any) -> bytes:
    core = document.core_properties
    core.author = "Strong Hire fixtures"
    core.title = "Resume"
    from datetime import datetime

    core.created = core.modified = datetime(2026, 1, 1)
    core.last_modified_by = "Strong Hire fixtures"
    core.revision = 1
    raw = io.BytesIO()
    document.save(raw)
    # Rewrite the zip with a fixed date on every entry so the file is the same on every build.
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw.getvalue())) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, FIXED_ZIP_DATE)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(fixed, src.read(info.filename))
    return out.getvalue()


def _bold_lead(paragraph: Any, lead: str, text: str) -> None:
    paragraph.add_run(lead + " ").bold = True
    paragraph.add_run(text)


def _docx_head(document: Any, case: Case) -> None:
    title = document.add_paragraph()
    run = title.add_run(case.person)
    run.bold = True
    run.font.size = Pt(16)
    document.add_paragraph(case.contact)
    if case.summary:
        document.add_heading("Summary", level=2)
        document.add_paragraph(case.summary)


def _docx_role_body(container: Any, role: Role, *, first: Any = None) -> None:
    """Role content into a document or a table cell. `first` reuses the cell's empty paragraph."""

    def para(style: str | None = None) -> Any:
        nonlocal first
        if first is not None:
            p, first = first, None
            if style:
                p.style = style
            return p
        return container.add_paragraph(style=style)

    if role.summary:
        para().add_run(role.summary)
    for sub in role.subs:
        _bold_lead(para(), sub.head, sub.text)
    for bullet in role.bullets:
        para("List Bullet").add_run(bullet)
    for heading, bullets in role.groups:
        para().add_run(heading).bold = True
        for bullet in bullets:
            para("List Bullet").add_run(bullet)
    if role.skills:
        para().add_run("Tools: " + ", ".join(role.skills))


def _docx_tail(document: Any, case: Case, *, skills: bool = True) -> None:
    if skills:
        document.add_heading("Skills", level=2)
        document.add_paragraph(", ".join(case.skills))
    document.add_heading("Education", level=2)
    for line in case.education_lines():
        document.add_paragraph(line)
    if case.certifications:
        document.add_heading("Certifications", level=2)
        for cert in case.certifications:
            document.add_paragraph(cert)


def docx_table(case: Case) -> bytes:
    """Experience as a two-column table (dates | details), then skills and education after it."""
    document = docx.Document()
    _docx_head(document, case)
    document.add_heading("Experience", level=2)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for role in case.roles:
        cells = table.add_row().cells
        cells[0].paragraphs[0].add_run(role.dates)
        head = cells[1].paragraphs[0]
        head.add_run(f"{role.company}, {role.place}").bold = True
        cells[1].add_paragraph(role.title)
        _docx_role_body(cells[1], role)
    _docx_tail(document, case)
    return _docx_bytes(document)


TEXTBOX_XML = """
<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:v="urn:schemas-microsoft-com:vml">
  <w:pict>
    <v:shape id="sidebar" type="#_x0000_t202" style="position:absolute;margin-left:380pt;
        margin-top:0;width:150pt;height:300pt;z-index:1">
      <v:textbox>
        <w:txbxContent>{paragraphs}</w:txbxContent>
      </v:textbox>
    </v:shape>
  </w:pict>
</w:r>
"""


def _xml_paragraph(text: str, bold: bool = False) -> str:
    from xml.sax.saxutils import escape

    props = "<w:rPr><w:b/></w:rPr>" if bold else ""
    return f'<w:p><w:r>{props}<w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'


def docx_textbox(case: Case) -> bytes:
    """The skills sit in a floating text box anchored in the first paragraph."""
    document = docx.Document()
    _docx_head(document, case)
    anchor = document.paragraphs[0]
    inner = _xml_paragraph("SKILLS", bold=True) + "".join(_xml_paragraph(s) for s in case.skills)
    anchor._p.append(parse_xml(TEXTBOX_XML.format(paragraphs=inner)))
    document.add_heading("Experience", level=2)
    for role in case.roles:
        p = document.add_paragraph()
        p.add_run(f"{role.title} - {role.company}, {role.place}").bold = True
        p.add_run().add_break(WD_BREAK.LINE)
        p.add_run(role.dates)
        _docx_role_body(document, role)
    _docx_tail(document, case, skills=False)
    return _docx_bytes(document)


def docx_subsections(case: Case) -> bytes:
    """Plain paragraphs; each role has bold lead-in sub-sections."""
    document = docx.Document()
    _docx_head(document, case)
    document.add_heading("Experience", level=2)
    for role in case.roles:
        p = document.add_paragraph()
        p.add_run(f"{role.company} | {role.place}").bold = True
        p.add_run("\t" + role.dates)
        document.add_paragraph().add_run(role.title).italic = True
        _docx_role_body(document, role)
    _docx_tail(document, case)
    return _docx_bytes(document)


# --- the cases ------------------------------------------------------------------------------

CASES: list[tuple[Case, str, Callable[[Case], bytes]]] = []


def case(ext: str, build: Callable[[Case], bytes]) -> Callable[[Callable[[], Case]], None]:
    def register(make: Callable[[], Case]) -> None:
        CASES.append((make(), ext, build))

    return register


@case("pdf", pdf_date_column)
def _date_column() -> Case:
    return Case(
        name="date-column",
        person="Rowan Ellery (synthetic person)",
        contact="rowan.ellery@example.com | 555-0141 | Bellevue, WA",
        summary="Engineering lead with 12 years in search, knowledge graphs and developer tools.",
        roles=[
            Role(
                company="Lumenfield Commerce",
                place="Bellevue, WA",
                title="Principal Engineer, AI Platform",
                dates="Sep 2025 \u2013 Present",
                start="2025-09",
                end=None,
                summary="Lead a group of 9 engineers that builds the AI platform for the "
                "marketplace search and seller tools.",
                subs=[
                    Sub(
                        "Knowledge Graphs & Catalog Search:",
                        "Designed a product knowledge "
                        "graph of 40 million items that raised search conversion by 3.1%.",
                    ),
                    Sub(
                        "AI Coding & Workflow Agent:",
                        "Led the prototype of an agentic platform "
                        "that drafts code changes; 120 engineers used it in the first month.",
                    ),
                    Sub(
                        "Evaluation Harness:",
                        "Built an offline evaluation suite of 2,400 "
                        "queries that gates every ranking model release.",
                    ),
                    Sub("Hiring:", "Interviewed 60 candidates and hired 5 senior engineers."),
                ],
            ),
            Role(
                company="Graniteview Cloud Works",
                place="Redmond, WA",
                title="Senior Software Engineer",
                dates="2017 \u2013 2025",
                start="2017",
                end="2025",
                summary="Owned the indexing service behind the developer documentation search.",
                subs=[
                    Sub(
                        "Indexing Pipeline:",
                        "Rebuilt the nightly index as a streaming job, "
                        "cutting freshness lag from 24 hours to 15 minutes.",
                    ),
                    Sub("Reliability:", "Reduced paging incidents from 14 to 3 per quarter."),
                ],
            ),
            Role(
                company="Bramblewood Labs",
                place="Seattle, WA",
                title="Software Engineer",
                dates="2013 \u2013 2017",
                start="2013",
                end="2017",
                summary="Built data tools for a 20-person research team.",
                subs=[
                    Sub("Query Builder:", "Shipped a visual query builder used by 300 analysts."),
                ],
            ),
        ],
        skills=["Python", "Java", "Elasticsearch", "Graph databases", "Kubernetes", "PyTorch"],
        education=[edu("Puget Ridge University", "M.S.", "Computer Science", "2013")],
    )


@case("pdf", pdf_sidebar)
def _sidebar() -> Case:
    return Case(
        name="sidebar-two-column",
        person="Imani Okafor-Lind (synthetic person)",
        contact="imani.ol@example.com | 555-0162 | Austin, TX",
        summary="Data engineer who builds batch and streaming pipelines for finance teams.",
        roles=[
            Role(
                company="Cobaltline Insurance",
                place="Austin, TX",
                title="Senior Data Engineer",
                dates="Jan 2022 - Present",
                start="2022-01",
                end=None,
                bullets=[
                    "Moved 140 nightly ETL jobs from cron to Airflow with retries, alerts "
                    "and an on-call runbook for the finance data team",
                    "Cut the monthly warehouse bill from $61,000 to $38,000 by clustering tables",
                    "Built a claims fraud feature store that 4 model teams read for "
                    "training and for real-time scoring of new claims",
                ],
            ),
            Role(
                company="Pebblestone Credit Union",
                place="Dallas, TX",
                title="Data Engineer",
                dates="Jun 2018 - Dec 2021",
                start="2018-06",
                end="2021-12",
                bullets=[
                    "Wrote the change-data-capture pipeline for 30 core banking tables",
                    "Added data quality checks on every load that caught 95% of bad loads "
                    "before the morning regulatory reports ran",
                ],
            ),
        ],
        skills=["SQL", "Python", "Airflow", "Snowflake", "dbt", "Kafka", "Terraform"],
        education=[edu("Hill Country State University", "B.S.", "Statistics", "2018")],
        certifications=["SnowPro Core Certification"],
    )


@case("pdf", pdf_single_column)
def _bold_subsections() -> Case:
    return Case(
        name="bold-subsections",
        person="Tomasz Varga (synthetic person)",
        contact="tomasz.varga@example.com | 555-0177 | Chicago, IL",
        summary="Mobile engineer focused on payments and app performance.",
        roles=[
            Role(
                company="Harborview Wallet",
                place="Chicago, IL",
                title="Staff iOS Engineer",
                dates="Mar 2021 - Present",
                start="2021-03",
                end=None,
                groups=[
                    (
                        "Payments",
                        [
                            "Rebuilt the card checkout flow; failed payments fell by 22%",
                            "Added Apple Pay for 1.2 million users in 6 weeks",
                        ],
                    ),
                    (
                        "Performance",
                        [
                            "Cut cold start time from 3.4 s to 1.1 s",
                            "Set up crash triage that took the crash-free rate to 99.8%",
                        ],
                    ),
                    ("Leadership", ["Mentored 6 engineers and ran the iOS guild"]),
                ],
                skills=["Swift", "SwiftUI"],
            ),
            Role(
                company="Gristmill Recipes",
                place="Remote",
                title="iOS Engineer",
                dates="Aug 2017 - Feb 2021",
                start="2017-08",
                end="2021-02",
                groups=[
                    (
                        "Features",
                        [
                            "Built offline recipe sync used by 400,000 monthly users",
                            "Launched the shopping list widget",
                        ],
                    ),
                ],
                skills=["Objective-C", "Swift"],
            ),
        ],
        skills=["Swift", "SwiftUI", "Objective-C", "Core Data", "XCTest", "Fastlane"],
        education=[edu("Lakeview College", "B.A.", "Mathematics", "2017")],
    )


@case("pdf", pdf_table)
def _table() -> Case:
    return Case(
        name="table-layout",
        person="Priya Ramanathan (synthetic person)",
        contact="priya.r@example.com | 555-0188 | Raleigh, NC",
        summary=None,
        roles=[
            Role(
                company="Quarry Point Health",
                place="Raleigh, NC",
                title="Technical Program Manager",
                dates="2023-04 - present",
                start="2023-04",
                end=None,
                bullets=[
                    "Ran the move of 3 clinics to the new scheduling system with zero downtime",
                    "Cut the release cycle from 6 weeks to 2 weeks across 5 teams",
                ],
            ),
            Role(
                company="Fenwick Logistics",
                place="Charlotte, NC",
                title="Program Manager",
                dates="2019-02 - 2023-03",
                start="2019-02",
                end="2023-03",
                bullets=[
                    "Delivered the warehouse scanner rollout to 41 sites on budget",
                    "Set up a risk review that closed 80% of audit findings in one quarter",
                ],
            ),
            Role(
                company="Tarheel Software Group",
                place="Durham, NC",
                title="Business Analyst",
                dates="2016-07 - 2019-01",
                start="2016-07",
                end="2019-01",
                bullets=["Wrote requirements for a billing portal used by 9,000 customers"],
            ),
        ],
        skills=["Jira", "SQL", "Risk management", "Agile", "Confluence"],
        education=[edu("Piedmont University", "B.S.", "Industrial Engineering", "2016")],
        certifications=["PMP"],
    )


def _multipage_pdf(case: Case) -> bytes:
    return pdf_single_column(case, footer=True, header=True, top=740)


@case("pdf", _multipage_pdf)
def _multipage() -> Case:
    def filler(n: int, topic: str) -> list[str]:
        return [
            f"{topic} project {i}: delivered milestone {i} with {10 + i} test cases"
            for i in range(1, n + 1)
        ]

    return Case(
        name="multipage-split-role",
        person="Desmond Achterberg (synthetic person)",
        contact="d.achterberg@example.com | 555-0199 | Denver, CO",
        summary="Backend engineer with 10 years of work on logistics and routing systems.",
        roles=[
            Role(
                company="Summitline Routing",
                place="Denver, CO",
                title="Lead Backend Engineer",
                dates="May 2022 - Present",
                start="2022-05",
                end=None,
                summary="Lead the routing platform team of 7 engineers.",
                bullets=[
                    "Rewrote the route optimizer in Rust; plans compute 5 times faster",
                    "Cut cloud spend by $410,000 a year with spot instances",
                    *filler(16, "Routing"),
                ],
            ),
            Role(
                company="Kestrel Freightworks",
                place="Boulder, CO",
                title="Senior Backend Engineer",
                dates="Feb 2018 - Apr 2022",
                start="2018-02",
                end="2022-04",
                bullets=[
                    "Built the shipment tracking API that serves 2,000 requests per second",
                    *filler(24, "Tracking"),
                    "Moved the event bus to Kafka without losing a single message",
                    "Led the on-call rotation for 18 months",
                ],
            ),
            Role(
                company="Alpenglow Maps",
                place="Golden, CO",
                title="Software Engineer",
                dates="Jul 2015 - Jan 2018",
                start="2015-07",
                end="2018-01",
                bullets=["Wrote the tile cache that cut map load time by 40%"],
            ),
        ],
        skills=["Rust", "Go", "PostgreSQL", "Kafka", "AWS"],
        education=[
            edu("Front Range Institute of Technology", "B.S.", "Computer Engineering", "2015")
        ],
    )


@case("pdf", pdf_same_company)
def _same_company() -> Case:
    return Case(
        name="same-company-two-roles",
        person="Hollis Brennan-Ito (synthetic person)",
        contact="hollis.bi@example.com | 555-0123 | Portland, OR",
        summary="Machine learning engineer who ships ranking and forecasting models.",
        roles=[
            Role(
                company="Northgate Freight",
                place="Portland, OR",
                title="Staff ML Engineer",
                dates="Jun 2023 - Present",
                start="2023-06",
                end=None,
                bullets=[
                    "Led the demand forecasting model that cut empty truck miles by 9%",
                    "Set the model review process for 3 ML teams",
                ],
            ),
            Role(
                company="Northgate Freight",
                place="Portland, OR",
                title="Senior ML Engineer",
                dates="Jan 2020 - May 2023",
                start="2020-01",
                end="2023-05",
                bullets=[
                    "Built the load matching ranker; accepted matches rose by 14%",
                    "Moved training from notebooks to a scheduled pipeline on Kubeflow",
                ],
            ),
            Role(
                company="Willowbend Analytics",
                place="Eugene, OR",
                title="Data Scientist",
                dates="Sep 2016 - Dec 2019",
                start="2016-09",
                end="2019-12",
                bullets=["Built churn models for 12 retail clients"],
            ),
        ],
        skills=["Python", "PyTorch", "XGBoost", "Kubeflow", "SQL"],
        education=[edu("Cascadia Polytechnic", "M.S.", "Statistics", "2016")],
    )


@case("docx", docx_table)
def _docx_table_case() -> Case:
    return Case(
        name="docx-table",
        person="Marisol Quintero (synthetic person)",
        contact="marisol.q@example.com | 555-0134 | Phoenix, AZ",
        summary="Product designer with 7 years in health and fintech apps.",
        roles=[
            Role(
                company="Saguaro Health",
                place="Phoenix, AZ",
                title="Senior Product Designer",
                dates="Oct 2022 - Present",
                start="2022-10",
                end=None,
                summary="Own the design of the patient app and its design system.",
                bullets=[
                    "Redesigned appointment booking; completed bookings rose by 27%",
                    "Built a design system of 80 components used by 4 squads",
                ],
            ),
            Role(
                company="Copperleaf Bank",
                place="Tempe, AZ",
                title="Product Designer",
                dates="Mar 2018 - Sep 2022",
                start="2018-03",
                end="2022-09",
                bullets=[
                    "Designed the mobile onboarding flow that cut drop-off by 18%",
                    "Ran 40 usability sessions a year with customers",
                ],
            ),
        ],
        skills=["Figma", "Prototyping", "User research", "Design systems", "Accessibility"],
        education=[edu("Sonoran School of Design", "B.F.A.", "Interaction Design", "2018")],
    )


@case("docx", docx_textbox)
def _docx_textbox_case() -> Case:
    return Case(
        name="docx-textbox-sidebar",
        person="Kwame Asante-Ross (synthetic person)",
        contact="kwame.ar@example.com | 555-0156 | Atlanta, GA",
        summary="Site reliability engineer for high-traffic consumer services.",
        roles=[
            Role(
                company="Peachtree Streaming",
                place="Atlanta, GA",
                title="Senior SRE",
                dates="Feb 2021 - Present",
                start="2021-02",
                end=None,
                bullets=[
                    "Took availability of the video API from 99.5% to 99.95%",
                    "Wrote the incident review process used by 11 teams",
                ],
            ),
            Role(
                company="Riverbend Ticketing",
                place="Savannah, GA",
                title="Systems Engineer",
                dates="May 2017 - Jan 2021",
                start="2017-05",
                end="2021-01",
                bullets=["Automated server builds with Ansible, from 2 days to 40 minutes"],
            ),
        ],
        skills=["Kubernetes", "Prometheus", "Terraform", "Go", "Ansible", "Linux"],
        education=[edu("Chattahoochee Tech University", "B.S.", "Information Systems", "2017")],
    )


@case("docx", docx_subsections)
def _docx_subsections_case() -> Case:
    return Case(
        name="docx-bold-subsections",
        person="Elodie Marchetti (synthetic person)",
        contact="elodie.m@example.com | 555-0110 | Boston, MA",
        summary="Product manager for B2B analytics products.",
        roles=[
            Role(
                company="Beacon Hill Analytics",
                place="Boston, MA",
                title="Senior Product Manager",
                dates="Apr 2022 - Present",
                start="2022-04",
                end=None,
                summary="Own the dashboards product line with $14M in annual revenue.",
                subs=[
                    Sub(
                        "Self-serve Dashboards:",
                        "Launched a template gallery that doubled weekly active creators to 5,200.",
                    ),
                    Sub("Pricing:", "Moved 600 accounts to usage-based pricing with 2% churn."),
                    Sub("Discovery:", "Ran 45 customer interviews that set the 2024 roadmap."),
                ],
            ),
            Role(
                company="Harborlight Software",
                place="Cambridge, MA",
                title="Product Manager",
                dates="Jul 2018 - Mar 2022",
                start="2018-07",
                end="2022-03",
                subs=[
                    Sub("Data Connectors:", "Shipped 12 new data connectors in one year."),
                    Sub("Onboarding:", "Cut time to first report from 3 days to 4 hours."),
                ],
            ),
        ],
        skills=["Product discovery", "SQL", "Amplitude", "Roadmapping", "A/B testing"],
        education=[edu("Charles River College", "B.A.", "Economics", "2018")],
    )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for item, ext, build in CASES:
        (OUT / f"{item.name}.{ext}").write_bytes(build(item))
        expected = json.dumps(item.expected(), indent=2, ensure_ascii=False) + "\n"
        (OUT / f"{item.name}.json").write_text(expected, encoding="utf-8")
        print(f"wrote {item.name}.{ext}")


if __name__ == "__main__":
    main()
